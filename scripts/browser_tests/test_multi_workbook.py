"""
Browser verification tests for Multi-Workbook Layout Support (All 6 APQC workbooks),
Benchmarking Metrics inspector UI, Level 6 rendering and filtering,
actionable error handling, and XSS escaping.
"""

import os
import time
import json
import base64
import openpyxl
from browser_tests.cdp_client import load_workbook

def run_multi_workbook_tests(client, project_dir):
    print("\n" + "="*70)
    print("PHASE: MULTI-WORKBOOK LAYOUTS, METRICS UI & LEVEL 6 VERIFICATION")
    print("="*70)

    fixture_path = os.path.join(project_dir, "data", "workbooks_expected.json")
    with open(fixture_path, "r", encoding="utf-8") as f:
        expected_fixture = json.load(f)

    # 1. Measure and Verify In-Browser Parsing for All 6 Workbooks
    print("\n--- 1. In-Browser Workbook Upload & Timing (All 6 Workbooks) ---")
    parse_timings = {}

    for key, exp in expected_fixture.items():
        wb_path = os.path.join(project_dir, exp["filename"])
        assert os.path.exists(wb_path), f"Missing {wb_path}"

        with open(wb_path, "rb") as f:
            b64_data = base64.b64encode(f.read()).decode("ascii")

        print(f"\n• Uploading & parsing {exp['name']} in Chrome:")
        t0 = time.time()
        client.evaluate(f"""(() => {{
            const binary = atob("{b64_data}");
            const bytes = new Uint8Array(binary.length);
            for (let i = 0; i < binary.length; i++) {{
                bytes[i] = binary.charCodeAt(i);
            }}
            pendingWorkbookBuffer = bytes;
            pendingFileName = "{exp['filename']}";
            const tStart = performance.now();
            openUploadModal(bytes, pendingFileName);
            window.__lastParseDurationMs = Math.round(performance.now() - tStart);
        }})()""")
        time.sleep(0.3)
        parse_ms = client.evaluate("window.__lastParseDurationMs")
        parse_timings[exp["name"]] = parse_ms
        print(f"  Parse time in browser: {parse_ms} ms")

        # Verify preview modal
        modal_visible = client.evaluate("document.getElementById('upload-modal-container').style.display !== 'none'")
        summary_text = client.evaluate("document.getElementById('modal-detected-summary').textContent.trim()")
        assert modal_visible, f"Modal did not open for {exp['name']}"
        assert f"{exp['elements']:,}" in summary_text or str(exp['elements']) in summary_text, (
            f"Expected {exp['elements']} elements in summary, got: {summary_text}"
        )
        assert "Detected layout:" in summary_text, f"Layout label missing in summary: {summary_text}"

        # Confirm upload
        client.evaluate("document.getElementById('confirm-upload-btn').click()")
        time.sleep(0.4)

        loaded_info = client.evaluate("""({
            elements: state.data.elements.length,
            layout: state.data.metadata.source_layout,
            definitions: state.data.metadata.definitions_count,
            metricRows: state.data.metadata.total_metric_records,
            maxLevel: Math.max(...state.data.elements.map(e => e.level))
        })""")

        print(f"  Elements: {loaded_info['elements']} (Expected: {exp['elements']})")
        print(f"  Layout Profile: {loaded_info['layout']} (Expected: {exp['layout_profile']})")
        expected_defs = exp.get("definitions_non_empty", exp["definitions_count"])
        print(f"  Definitions: {loaded_info['definitions']} (Expected: {expected_defs})")
        print(f"  Metric Records: {loaded_info['metricRows']} (Expected: {exp['metric_rows']})")
        print(f"  Max Hierarchy Level: {loaded_info['maxLevel']} (Expected: {max(exp['levels'])})")

        assert loaded_info["elements"] == exp["elements"], f"{exp['name']}: Elements mismatch"
        assert loaded_info["layout"] == exp["layout_profile"], f"{exp['name']}: Layout mismatch"
        assert loaded_info["definitions"] == expected_defs, f"{exp['name']}: Definitions mismatch"
        assert loaded_info["metricRows"] == exp["metric_rows"], f"{exp['name']}: Metric rows mismatch"
        assert loaded_info["maxLevel"] == max(exp["levels"]), f"{exp['name']}: Max level mismatch"

        # Verify visible tabs (all 4 canonical screens should be shown for these 6 workbooks)
        tabs = client.evaluate("Array.from(document.querySelectorAll('#view-tabs-nav .view-tab-btn')).map(b => b.getAttribute('data-view'))")
        assert set(tabs) == {"explorer", "matrix", "analytics", "glossary"}, f"{exp['name']}: Visible tabs mismatch: {tabs}"
        print(f"  ✓ {exp['name']} browser upload, parse & tabs verified.")

    # 1B. Family B Glossary Verification (Definitions from Element Description)
    print("\n--- 1B. Family B Glossary Verification (Definitions from Element Description) ---")
    for b_key in ["cross_industry", "consumer_products"]:
        b_exp = expected_fixture[b_key]
        b_path = os.path.join(project_dir, b_exp["filename"])
        load_workbook(client, b_path)

        # Switch to Glossary view
        client.evaluate("switchView('glossary')")
        time.sleep(0.3)
        ui_glossary_count = client.evaluate("""(() => {
            const countEl = document.getElementById('glossary-subtitle');
            return countEl ? parseInt(countEl.textContent.replace(/[^0-9]/g, ''), 10) : 0;
        })()""")
        print(f"  • {b_exp['name']} Glossary UI Count: {ui_glossary_count} (Expected: {b_exp['definitions_non_empty']})")
        assert ui_glossary_count == b_exp["definitions_non_empty"], f"Glossary count mismatch for {b_exp['name']}"

        # Verify sampled elements definitions match Element Description from raw workbook
        wb_b = openpyxl.load_workbook(b_path, data_only=True)
        ws_comb = wb_b['Combined']
        headers = [str(c.value).strip().lower() if c.value else "" for c in ws_comb[1]]
        hid_col = headers.index("hierarchy id") + 1
        desc_col = headers.index("element description") + 1

        samples = ["1.1", "1.1.1", "1.1.1.1"]
        raw_defs = {}
        for r in range(2, min(500, ws_comb.max_row + 1)):
            hid_val = str(ws_comb.cell(row=r, column=hid_col).value or "").strip()
            if hid_val in samples and hid_val not in raw_defs:
                desc_val = str(ws_comb.cell(row=r, column=desc_col).value or "").strip()
                raw_defs[hid_val] = desc_val
                if len(raw_defs) == len(samples):
                    break

        for hid, expected_desc in raw_defs.items():
            parsed_def = client.evaluate(f"state.data.elements.find(e => e.hierarchy_id === '{hid}')?.definition")
            print(f"    - Sample {hid}: '{parsed_def[:50]}...' matches raw")
            assert parsed_def == expected_desc, f"Definition mismatch for {hid} in {b_exp['name']}"
    print("  ✓ Family B Glossary definitions verified against Element Description.")

    # 2. Benchmarking Metrics Inspector UI Verification (Family B: Cross-Industry)
    print("\n--- 2. Benchmarking Metrics Inspector UI (Cross-Industry v8.0) ---")
    # Load Cross-Industry workbook again to test metrics
    ci_path = os.path.join(project_dir, expected_fixture["cross_industry"]["filename"])
    with open(ci_path, "rb") as f:
        ci_b64 = base64.b64encode(f.read()).decode("ascii")

    client.evaluate(f"""(() => {{
        const binary = atob("{ci_b64}");
        const bytes = new Uint8Array(binary.length);
        for (let i = 0; i < binary.length; i++) {{
            bytes[i] = binary.charCodeAt(i);
        }}
        pendingWorkbookBuffer = bytes;
        pendingFileName = "{expected_fixture['cross_industry']['filename']}";
        openUploadModal(bytes, pendingFileName);
    }})()""")
    time.sleep(0.3)
    client.evaluate("document.getElementById('confirm-upload-btn').click()")
    time.sleep(0.4)

    # Find element WITH metrics and inspect
    metric_elem_info = client.evaluate("""(() => {
        const elemWithMetrics = state.data.elements.find(e => e.metrics && e.metrics.length > 0);
        if (!elemWithMetrics) return null;
        inspectElement(elemWithMetrics, false);
        const drawer = document.getElementById("inspector-content");
        const metricsSection = drawer.querySelector("details.ref-section summary span");
        const metricCards = drawer.querySelectorAll(".metric-item-card");
        const firstCard = metricCards[0];
        return {
            id: elemWithMetrics.id,
            hid: elemWithMetrics.hierarchy_id,
            metricsCount: elemWithMetrics.metrics.length,
            sectionSummaryText: metricsSection ? metricsSection.textContent.trim() : "",
            cardCount: metricCards.length,
            firstCardText: firstCard ? firstCard.textContent.trim() : ""
        };
    })()""")
    print(f"Element with metrics: HID {metric_elem_info['hid']} ({metric_elem_info['metricsCount']} metrics)")
    print(f"Inspector section header: '{metric_elem_info['sectionSummaryText']}'")
    print(f"Rendered metric cards: {metric_elem_info['cardCount']}")
    print(f"Sample metric card text: '{metric_elem_info['firstCardText'][:70]}...'")

    assert f"Benchmarking metrics ({metric_elem_info['metricsCount']})" in metric_elem_info["sectionSummaryText"], (
        "Inspector does not display correct Benchmarking metrics count header"
    )
    assert metric_elem_info["cardCount"] == metric_elem_info["metricsCount"], (
        f"Expected {metric_elem_info['metricsCount']} metric cards, got {metric_elem_info['cardCount']}"
    )

    # Find element WITHOUT metrics and inspect
    no_metric_elem_info = client.evaluate("""(() => {
        const elemNoMetrics = state.data.elements.find(e => !e.metrics || e.metrics.length === 0);
        if (!elemNoMetrics) return null;
        inspectElement(elemNoMetrics, false);
        const drawer = document.getElementById("inspector-content");
        const hasMetricsHeader = drawer.textContent.includes("Benchmarking metrics");
        return {
            id: elemNoMetrics.id,
            hid: elemNoMetrics.hierarchy_id,
            hasMetricsHeader: hasMetricsHeader
        };
    })()""")
    print(f"Element without metrics: HID {no_metric_elem_info['hid']}")
    print(f"Inspector contains 'Benchmarking metrics': {no_metric_elem_info['hasMetricsHeader']} (Expected: False)")
    assert not no_metric_elem_info["hasMetricsHeader"], "Benchmarking metrics section must be omitted when empty"
    print("  ✓ Benchmarking metrics UI presence and omission verified.")

    # 3. Level 6 Hierarchy Verification (Health Insurance Payor)
    print("\n--- 3. Level 6 Hierarchy Verification (Health Insurance Payor) ---")
    payor_path = os.path.join(project_dir, expected_fixture["payor"]["filename"])
    with open(payor_path, "rb") as f:
        payor_b64 = base64.b64encode(f.read()).decode("ascii")

    client.evaluate(f"""(() => {{
        const binary = atob("{payor_b64}");
        const bytes = new Uint8Array(binary.length);
        for (let i = 0; i < binary.length; i++) {{
            bytes[i] = binary.charCodeAt(i);
        }}
        pendingWorkbookBuffer = bytes;
        pendingFileName = "{expected_fixture['payor']['filename']}";
        openUploadModal(bytes, pendingFileName);
    }})()""")
    time.sleep(0.3)
    client.evaluate("document.getElementById('confirm-upload-btn').click()")
    time.sleep(0.4)

    # Check dynamic Level Filter options
    filter_options = client.evaluate("""(() => {
        const select = document.getElementById("level-filter");
        return Array.from(select.options).map(o => ({ value: o.value, text: o.textContent.trim() }));
    })()""")
    has_level_6 = any(o["value"] == "6" for o in filter_options)
    print(f"Level filter options count: {len(filter_options)}")
    print(f"Level 6 option exists: {has_level_6}")
    assert has_level_6, "Dynamic level filter missing Level 6 option"

    # Select all categories so all level 6 elements across Categories 2 and 3 are in scope
    client.evaluate("selectCategory('all')")
    time.sleep(0.2)

    # Filter by level 6 and verify visible count across all categories
    client.evaluate("""(() => {
        const select = document.getElementById("level-filter");
        select.value = "6";
        select.dispatchEvent(new Event("change"));
    })()""")
    time.sleep(0.2)
    level_6_all_text = client.evaluate("document.getElementById('visible-count').textContent.trim()")
    print(f"Visible elements across all categories filtered to Level 6: {level_6_all_text} (Expected: 13)")
    assert level_6_all_text == "13", f"Expected 13 level-6 elements across all categories, got {level_6_all_text}"

    # Select Category 2 and verify visible count in Category 2
    client.evaluate("selectCategory('2')")
    time.sleep(0.2)
    level_6_cat2_text = client.evaluate("document.getElementById('visible-count').textContent.trim()")
    print(f"Visible elements in Category 2 filtered to Level 6: {level_6_cat2_text} (Expected: 10)")
    assert level_6_cat2_text == "10", f"Expected 10 level-6 elements in Category 2, got {level_6_cat2_text}"

    # Reset level filter to all
    client.evaluate("""(() => {
        const select = document.getElementById("level-filter");
        select.value = "all";
        select.dispatchEvent(new Event("change"));
    })()""")
    time.sleep(0.2)

    # Verify Level 6 element rendering, indentation, and breadcrumbs
    l6_inspect_info = client.evaluate("""(() => {
        const l6Elem = state.data.elements.find(e => e.level === 6);
        if (!l6Elem) return null;
        jumpToElement(l6Elem.id);
        const row = document.querySelector(`.tree-row[data-id="${l6Elem.id}"]`);
        const drawer = document.getElementById("inspector-content");
        const breadcrumb = drawer.querySelector(".ref-breadcrumb");
        const subline = drawer.querySelector(".ref-subline");
        return {
            id: l6Elem.id,
            hid: l6Elem.hierarchy_id,
            paddingLeft: row ? row.style.paddingLeft : "",
            breadcrumbText: breadcrumb ? breadcrumb.textContent.trim() : "",
            sublineText: subline ? subline.textContent.trim() : ""
        };
    })()""")
    print(f"Level 6 element: HID {l6_inspect_info['hid']}")
    print(f"Row padding-left: {l6_inspect_info['paddingLeft']} (Expected: 5.375rem)")
    print(f"Inspector Subline: {l6_inspect_info['sublineText']}")
    print(f"Breadcrumb: {l6_inspect_info['breadcrumbText']}")
    assert l6_inspect_info["paddingLeft"] == "5.375rem", f"Unexpected padding: {l6_inspect_info['paddingLeft']}"
    assert "Level 6" in l6_inspect_info["sublineText"] or "Sub-task" in l6_inspect_info["sublineText"]
    print("  ✓ Level 6 hierarchy, filter, indentation, and breadcrumbs verified.")

    # 4. Actionable Error Handling in UI
    print("\n--- 4. Actionable Error Handling in UI ---")
    error_info = client.evaluate("""(() => {
        try {
            const dummyWb = {
                SheetNames: ["Introduction", "Combined"],
                Sheets: {
                    Introduction: {},
                    Combined: XLSX.utils.aoa_to_sheet([["Col A", "Col B"], [1, 2]])
                }
            };
            openUploadModal(dummyWb, "unknown.xlsx");
        } catch (e) {}
        const errBox = document.getElementById("upload-dialog-error");
        return {
            visible: errBox ? errBox.style.display !== "none" : false,
            text: errBox ? errBox.textContent : ""
        };
    })()""")
    print(f"Error box visible: {error_info['visible']}")
    print(f"Error message snippet:\n    {error_info['text'].split(chr(10))[0]}")
    assert error_info["visible"], "Error box should be visible on unknown layout"
    assert "Unknown workbook layout" in error_info["text"], "Actionable error missing layout info"
    assert "Expected either:" in error_info["text"], "Actionable error missing guidance"
    client.evaluate("closeUploadModal()")
    print("  ✓ Actionable layout error correctly shown in upload modal.")

    # 5. Escaping & Security for Metrics and Element Description
    print("\n--- 5. XSS Escaping for Metrics and Element Descriptions ---")
    xss_test_result = client.evaluate("""(() => {
        window.__xss_fired = false;
        const fakeElem = {
            id: "xss_test_elem",
            pcf_id: "99999",
            hierarchy_id: "1.9.9",
            name: "XSS Test Process <script>window.__xss_fired=true;</script>",
            level: 3,
            level_name: "Process",
            category_id: "1",
            parent_hierarchy_id: "1.0",
            diff_index: 0,
            metrics_available: true,
            process_type: "Operating Processes",
            definition: "Description with <img src=x onerror=window.__xss_fired=true> payload",
            change_details: "",
            has_changes: false,
            is_new: false,
            is_renamed: false,
            renamed_was: "",
            added_children: [],
            removed_children: [],
            changed_children: [],
            metrics: [
                {
                    category: "Cat <script>window.__xss_fired=true;</script>",
                    id: "M999",
                    name: "Metric <img src=x onerror=window.__xss_fired=true>",
                    formula: "A / B <script>window.__xss_fired=true;</script>",
                    units: "USD <b onmouseover=window.__xss_fired=true>bold</b>"
                }
            ]
        };
        inspectElement(fakeElem, false);
        return {
            xssFired: Boolean(window.__xss_fired),
            inspectorHtml: document.getElementById("inspector-content").innerHTML
        };
    })()""")
    print(f"XSS payload executed: {xss_test_result['xssFired']} (Expected: False)")
    assert not xss_test_result["xssFired"], "XSS payload executed in inspector!"
    assert "<script>" not in xss_test_result["inspectorHtml"], "Unescaped <script> tag in DOM"
    assert "<img src=x" not in xss_test_result["inspectorHtml"], "Unescaped <img> tag in DOM"
    print("  ✓ Metric & description XSS sanitization verified.")

    # 6. Restore Healthcare PCF
    print("\n--- 6. Restoring Healthcare PCF ---")
    healthcare_path = os.path.join(project_dir, expected_fixture["healthcare"]["filename"])
    load_workbook(client, healthcare_path)
    restored_count = client.evaluate("state.data ? state.data.elements.length : 0")
    print(f"Restored dataset elements: {restored_count} (Expected: 1985)")
    assert restored_count == 1985, "Failed to restore Healthcare PCF"
    print("  ✓ State cleanly restored.")

    print("\n======================================================================")
    print("ALL MULTI-WORKBOOK AND METRIC UI TESTS PASSED SUCCESSFULLY!")
    print("Browser Parse Timings Summary:")
    for name, ms in parse_timings.items():
        print(f"  • {name}: {ms} ms")
    print("======================================================================")
