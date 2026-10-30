"""
Browser tests for App Startup (Empty state), View Registry,
Capabilities computation, State Machine transitions, Hiding unavailable screens,
and Actionable Failure recovery.
"""

import os
import time
import json
import base64
import openpyxl
import shutil
from browser_tests.cdp_client import load_workbook, close_workbook

def run_empty_start_tests(client, project_dir):
    print("\n" + "="*70)
    print("PHASE 0: EMPTY START, STATE MACHINE, CAPABILITIES & VIEW REGISTRY")
    print("="*70)

    # 1. Fresh page load in Empty State
    print("\n--- 1. Verifying Clean Empty Startup ---")
    empty_status = client.evaluate("""({
        appState: state.appState,
        data: state.data,
        capabilities: state.capabilities,
        emptyScreenVisible: window.getComputedStyle(document.getElementById('empty-state-screen')).display !== 'none',
        workspaceVisible: window.getComputedStyle(document.getElementById('workspace-container')).display !== 'none',
        closeBtnVisible: window.getComputedStyle(document.getElementById('close-workbook-btn')).display !== 'none',
        headerTabsVisible: window.getComputedStyle(document.getElementById('view-tabs-nav')).display !== 'none',
        searchVisible: (() => { const el = document.querySelector('.header-search'); return el ? window.getComputedStyle(el).display !== 'none' : false; })(),
        sidebarToggleVisible: window.getComputedStyle(document.getElementById('toggle-sidebar-btn')).display !== 'none',
        themeBtnVisible: (() => { const el = document.getElementById('theme-toggle-btn'); return el ? window.getComputedStyle(el).display !== 'none' : false; })(),
        textScaleVisible: (() => { const el = document.querySelector('.text-scale-group'); return el ? window.getComputedStyle(el).display !== 'none' : false; })(),
        dropZoneTabIndex: document.getElementById('empty-drop-zone')?.getAttribute('tabindex'),
        emptyLoadBtnExists: Boolean(document.getElementById('empty-load-btn')),
        resetBtnExists: Boolean(document.getElementById('reset-default-pcf-btn'))
    })""")

    print(f"  • App State: {empty_status['appState']} (Expected: Empty)")
    print(f"  • state.data is null: {empty_status['data'] is None}")
    print(f"  • state.capabilities is null: {empty_status['capabilities'] is None}")
    print(f"  • Empty screen visible: {empty_status['emptyScreenVisible']}")
    print(f"  • Workspace hidden: {not empty_status['workspaceVisible']}")
    print(f"  • Header controls hidden: {not empty_status['headerTabsVisible'] and not empty_status['searchVisible']}")
    print(f"  • Theme & Scale controls visible: {empty_status['themeBtnVisible'] and empty_status['textScaleVisible']}")
    print(f"  • Drop zone tabindex: {empty_status['dropZoneTabIndex']} (Expected: 0)")
    print(f"  • Reset-to-sample button removed: {not empty_status['resetBtnExists']}")

    assert empty_status['appState'] == "Empty", f"Expected appState 'Empty', got {empty_status['appState']}"
    assert empty_status['data'] is None, "state.data should be null on empty start"
    assert empty_status['capabilities'] is None, "state.capabilities should be null on empty start"
    assert empty_status['emptyScreenVisible'], "Empty state screen is not visible"
    assert not empty_status['workspaceVisible'], "Workspace should be hidden on empty start"
    assert not empty_status['headerTabsVisible'], "Header tabs should be hidden on empty start"
    assert not empty_status['searchVisible'], "Header search should be hidden on empty start"
    assert not empty_status['sidebarToggleVisible'], "Sidebar toggle should be hidden on empty start"
    assert empty_status['themeBtnVisible'], "Theme button should stay visible on empty start"
    assert empty_status['textScaleVisible'], "Text scale controls should stay visible on empty start"
    assert empty_status['dropZoneTabIndex'] == "0", "Drop zone must have tabindex 0 for keyboard accessibility"
    assert empty_status['emptyLoadBtnExists'], "Empty state load button missing"
    assert not empty_status['resetBtnExists'], "Reset-to-sample button should be deleted"

    # Verify Network requests: 0 requests for data/pcf_data.*
    pcf_data_requests = [req for req in client.network_requests if "data/pcf_data" in req]
    print(f"  • Network requests matching 'data/pcf_data.*': {len(pcf_data_requests)}")
    assert len(pcf_data_requests) == 0, f"Found unexpected network requests for bundled sample: {pcf_data_requests}"

    # Verify Theme and Scale work smoothly in Empty state
    theme_scale_empty = client.evaluate("""(() => {
        applyTheme('indigo-rose');
        stepTextScale(1);
        const th = document.documentElement.getAttribute('data-theme');
        const sc = state.textScale;
        applyTheme('dark');
        stepTextScale(-1);
        return { th, sc };
    })()""")
    print(f"  • Theme & scale work while empty: theme={theme_scale_empty['th']}, scale={theme_scale_empty['sc']}")
    assert theme_scale_empty['th'] == "indigo-rose", "Theme failed to apply in empty state"
    assert theme_scale_empty['sc'] == 1.15, "Scale failed to step up in empty state"

    # Verify Blocked Storage (SecurityError) during Empty Start
    storage_blocked = client.evaluate("""(() => {
        const origGet = Storage.prototype.getItem;
        Storage.prototype.getItem = function() { throw new DOMException("Storage blocked", "SecurityError"); };
        let ok = false;
        try {
            initTheme();
            ok = (state.appState === "Empty");
        } catch (e) {
            ok = false;
        } finally {
            Storage.prototype.getItem = origGet;
        }
        return { ok };
    })()""")
    print(f"  • Empty state survives blocked localStorage: {storage_blocked['ok']}")
    assert storage_blocked['ok'], "App crashed on empty start with blocked storage"

    # 2. Failure Cases from Empty State
    print("\n--- 2. Actionable Failure Recovery from Empty State ---")
    # A. Corrupt file from Empty
    corrupt_from_empty = client.evaluate("""(() => {
        const junk = new Uint8Array([0x50, 0x4b, 0x03, 0x04, 0xff, 0xff, 0xff, 0xff]);
        openUploadModal(junk, "corrupt.xlsx");
        const errBox = document.getElementById('upload-dialog-error');
        const res = {
            appState: state.appState,
            errVisible: errBox && errBox.style.display !== 'none',
            errText: errBox ? errBox.textContent : ''
        };
        closeUploadModal();
        return res;
    })()""")
    print(f"  • Corrupt file from Empty: state={corrupt_from_empty['appState']}, error='{corrupt_from_empty['errText'][:60]}...'")
    assert corrupt_from_empty['appState'] == "Empty", "Corrupt file changed appState away from Empty"
    assert corrupt_from_empty['errVisible'], "Error message not shown on corrupt file"
    assert "Cannot read workbook structure" in corrupt_from_empty['errText'] or "workbook" in corrupt_from_empty['errText'].lower()

    # B. Missing 'Combined' sheet from Empty
    missing_combined_xlsx = "/tmp/missing_combined.xlsx"
    wb_no_comb = openpyxl.Workbook()
    wb_no_comb.active.title = "Introduction"
    wb_no_comb.create_sheet("Categories")
    wb_no_comb.save(missing_combined_xlsx)
    with open(missing_combined_xlsx, "rb") as f:
        missing_b64 = base64.b64encode(f.read()).decode("ascii")

    missing_from_empty = client.evaluate(f"""(() => {{
        const bin = atob("{missing_b64}");
        const bytes = new Uint8Array(bin.length);
        for (let i = 0; i < bin.length; i++) bytes[i] = bin.charCodeAt(i);
        openUploadModal(bytes, "missing_combined.xlsx");
        const errBox = document.getElementById('upload-dialog-error');
        const res = {{
            appState: state.appState,
            errVisible: errBox && errBox.style.display !== 'none',
            errText: errBox ? errBox.textContent : ''
        }};
        closeUploadModal();
        return res;
    }})()""")
    print(f"  • Missing Combined sheet from Empty: state={missing_from_empty['appState']}, error='{missing_from_empty['errText']}'")
    assert missing_from_empty['appState'] == "Empty", "Missing Combined changed appState away from Empty"
    assert missing_from_empty['errVisible'], "Error message not shown on missing Combined sheet"
    assert "Combined" in missing_from_empty['errText'], "Error text should mention missing 'Combined' sheet"

    # C. Unknown layout from Empty
    unknown_layout_xlsx = "/tmp/unknown_layout.xlsx"
    wb_unk = openpyxl.Workbook()
    ws_comb = wb_unk.active
    ws_comb.title = "Combined"
    ws_comb.append(["Unrecognized A", "Unrecognized B", "Unrecognized C"])
    ws_comb.append([1, 2, 3])
    wb_unk.save(unknown_layout_xlsx)
    with open(unknown_layout_xlsx, "rb") as f:
        unk_b64 = base64.b64encode(f.read()).decode("ascii")

    unk_from_empty = client.evaluate(f"""(() => {{
        const bin = atob("{unk_b64}");
        const bytes = new Uint8Array(bin.length);
        for (let i = 0; i < bin.length; i++) bytes[i] = bin.charCodeAt(i);
        openUploadModal(bytes, "unknown_layout.xlsx");
        const errBox = document.getElementById('upload-dialog-error');
        const res = {{
            appState: state.appState,
            errVisible: errBox && errBox.style.display !== 'none',
            errText: errBox ? errBox.textContent : ''
        }};
        closeUploadModal();
        return res;
    }})()""")
    print(f"  • Unknown layout from Empty: state={unk_from_empty['appState']}, error snippet='{unk_from_empty['errText'].splitlines()[0]}'")
    assert unk_from_empty['appState'] == "Empty", "Unknown layout changed appState away from Empty"
    assert unk_from_empty['errVisible'], "Error message not shown on unknown layout"
    assert "Unknown workbook layout" in unk_from_empty['errText'], "Error text should explain unknown layout"
    assert "Expected either:" in unk_from_empty['errText'], "Error text should include actionable expected layouts"

    # 3. Successful Load: Empty -> Loading -> Loaded
    print("\n--- 3. Successful Load (Healthcare Provider) ---")
    healthcare_path = os.path.join(project_dir, "K07128_Healthcare Provider_v721_vs_Healthcare Provider_v611_April 2023.xlsx")
    load_workbook(client, healthcare_path, via_drop=True)

    loaded_status = client.evaluate("""({
        appState: state.appState,
        elementsCount: state.data ? state.data.elements.length : 0,
        capabilities: state.capabilities,
        emptyScreenVisible: window.getComputedStyle(document.getElementById('empty-state-screen')).display !== 'none',
        workspaceVisible: window.getComputedStyle(document.getElementById('workspace-container')).display !== 'none',
        closeBtnVisible: window.getComputedStyle(document.getElementById('close-workbook-btn')).display !== 'none',
        headerTitle: document.getElementById('header-framework-title')?.textContent.trim(),
        availableTabs: Array.from(document.querySelectorAll('#view-tabs-nav .view-tab-btn')).map(b => b.getAttribute('data-view'))
    })""")
    print(f"  • App State: {loaded_status['appState']} (Expected: Loaded)")
    print(f"  • Elements: {loaded_status['elementsCount']} (Expected: 1985)")
    print(f"  • Capabilities: {loaded_status['capabilities']}")
    print(f"  • Workspace visible: {loaded_status['workspaceVisible']}")
    print(f"  • Empty screen hidden: {not loaded_status['emptyScreenVisible']}")
    print(f"  • Close button visible: {loaded_status['closeBtnVisible']}")
    print(f"  • Header title: '{loaded_status['headerTitle']}'")
    print(f"  • Available view tabs: {loaded_status['availableTabs']}")

    assert loaded_status['appState'] == "Loaded", "App failed to transition to Loaded"
    assert loaded_status['elementsCount'] == 1985, f"Expected 1985 elements, got {loaded_status['elementsCount']}"
    assert loaded_status['workspaceVisible'], "Workspace should be visible when Loaded"
    assert not loaded_status['emptyScreenVisible'], "Empty screen should be hidden when Loaded"
    assert loaded_status['closeBtnVisible'], "Close button should be visible when Loaded"
    assert loaded_status['capabilities']['changes'] is True, "Healthcare should have changes=True"
    assert loaded_status['capabilities']['definitions'] > 0, "Healthcare should have definitions > 0"
    assert set(loaded_status['availableTabs']) == {"explorer", "matrix", "analytics", "glossary"}, "Expected all 4 canonical tabs for Healthcare"

    # 4. Failure Cases from Loaded State (Previous Workbook Must Stay Intact!)
    print("\n--- 4. Actionable Failure Recovery from Loaded State (Intact Previous State) ---")
    # A. Corrupt file while Loaded
    corrupt_from_loaded = client.evaluate("""(() => {
        const junk = new Uint8Array([0x50, 0x4b, 0x03, 0x04, 0xff, 0xff, 0xff, 0xff]);
        openUploadModal(junk, "corrupt2.xlsx");
        const errBox = document.getElementById('upload-dialog-error');
        const res = {
            appState: state.appState,
            elementsCount: state.data ? state.data.elements.length : 0,
            errVisible: errBox && errBox.style.display !== 'none',
            errText: errBox ? errBox.textContent : ''
        };
        closeUploadModal();
        return res;
    })()""")
    print(f"  • Corrupt file from Loaded: state={corrupt_from_loaded['appState']}, elements={corrupt_from_loaded['elementsCount']} (Expected: 1985)")
    assert corrupt_from_loaded['appState'] == "Loaded", "State should stay Loaded with previous data on corrupt upload"
    assert corrupt_from_loaded['elementsCount'] == 1985, "Previous dataset was modified or destroyed by corrupt file"
    assert corrupt_from_loaded['errVisible'], "Error message was not shown for corrupt file while Loaded"

    # B. Missing Combined sheet while Loaded
    missing_from_loaded = client.evaluate(f"""(() => {{
        const bin = atob("{missing_b64}");
        const bytes = new Uint8Array(bin.length);
        for (let i = 0; i < bin.length; i++) bytes[i] = bin.charCodeAt(i);
        openUploadModal(bytes, "missing_combined.xlsx");
        const errBox = document.getElementById('upload-dialog-error');
        const res = {{
            appState: state.appState,
            elementsCount: state.data ? state.data.elements.length : 0,
            errVisible: errBox && errBox.style.display !== 'none'
        }};
        closeUploadModal();
        return res;
    }})()""")
    print(f"  • Missing Combined from Loaded: state={missing_from_loaded['appState']}, elements={missing_from_loaded['elementsCount']} (Expected: 1985)")
    assert missing_from_loaded['appState'] == "Loaded", "State should stay Loaded on missing Combined"
    assert missing_from_loaded['elementsCount'] == 1985, "Previous dataset was lost on missing Combined"

    # C. Unknown layout while Loaded
    unk_from_loaded = client.evaluate(f"""(() => {{
        const bin = atob("{unk_b64}");
        const bytes = new Uint8Array(bin.length);
        for (let i = 0; i < bin.length; i++) bytes[i] = bin.charCodeAt(i);
        openUploadModal(bytes, "unknown_layout.xlsx");
        const errBox = document.getElementById('upload-dialog-error');
        const res = {{
            appState: state.appState,
            elementsCount: state.data ? state.data.elements.length : 0,
            errVisible: errBox && errBox.style.display !== 'none'
        }};
        closeUploadModal();
        return res;
    }})()""")
    print(f"  • Unknown layout from Loaded: state={unk_from_loaded['appState']}, elements={unk_from_loaded['elementsCount']} (Expected: 1985)")
    assert unk_from_loaded['appState'] == "Loaded", "State should stay Loaded on unknown layout"
    assert unk_from_loaded['elementsCount'] == 1985, "Previous dataset was lost on unknown layout"

    # 5. Loaded -> Empty via "Close workbook" action
    print("\n--- 5. Close Workbook Action (Loaded -> Empty) ---")
    close_workbook(client)
    closed_status = client.evaluate("""({
        appState: state.appState,
        data: state.data,
        capabilities: state.capabilities,
        emptyScreenVisible: window.getComputedStyle(document.getElementById('empty-state-screen')).display !== 'none',
        workspaceVisible: window.getComputedStyle(document.getElementById('workspace-container')).display !== 'none'
    })""")
    print(f"  • App State after close: {closed_status['appState']} (Expected: Empty)")
    print(f"  • Empty screen visible: {closed_status['emptyScreenVisible']}")
    print(f"  • Workspace hidden: {not closed_status['workspaceVisible']}")
    assert closed_status['appState'] == "Empty", "Close button failed to transition state to Empty"
    assert closed_status['data'] is None, "state.data should be null after closing workbook"
    assert closed_status['emptyScreenVisible'], "Empty screen should be visible after closing workbook"
    assert not closed_status['workspaceVisible'], "Workspace should be hidden after closing workbook"

    # 6. Screen Hiding & Capabilities Fallback
    # Test workbook with NO definitions and NO changes
    print("\n--- 6. Screen Hiding & Fallback (0 definitions & 0 changes) ---")
    no_defs_no_changes_xlsx = "/tmp/no_defs_no_changes.xlsx"
    shutil.copyfile(healthcare_path, no_defs_no_changes_xlsx)
    wb_hide = openpyxl.load_workbook(no_defs_no_changes_xlsx)
    if "Glossary terms" in wb_hide.sheetnames:
        ws_g = wb_hide["Glossary terms"]
        ws_g.delete_rows(2, ws_g.max_row)
    ws_comb = wb_hide["Combined"]
    headers = [str(c.value).strip().lower() if c.value else "" for c in ws_comb[1]]
    chg_col = headers.index("change details") + 1
    for r in range(2, ws_comb.max_row + 1):
        ws_comb.cell(row=r, column=chg_col, value="")
    wb_hide.save(no_defs_no_changes_xlsx)

    # First load full Healthcare workbook and switch to Glossary view
    load_workbook(client, healthcare_path)
    client.evaluate("switchView('glossary')")
    time.sleep(0.3)
    curr_view = client.evaluate("state.activeView")
    print(f"  • Loaded full workbook and switched to: {curr_view}")
    assert curr_view == "glossary", "Failed to switch to glossary view"

    # Now load the workbook with 0 definitions and 0 changes
    load_workbook(client, no_defs_no_changes_xlsx)
    time.sleep(0.3)

    hiding_status = client.evaluate("""({
        activeView: state.activeView,
        changesCap: state.capabilities.changes,
        defsCap: state.capabilities.definitions,
        matrixTabExists: Boolean(document.getElementById('view-tab-matrix')),
        glossaryTabExists: Boolean(document.getElementById('view-tab-glossary')),
        explorerTabExists: Boolean(document.getElementById('view-tab-explorer')),
        analyticsTabExists: Boolean(document.getElementById('view-tab-analytics')),
        disabledTabsCount: document.querySelectorAll('#view-tabs-nav .view-tab-btn[disabled]').length,
        availableTabs: Array.from(document.querySelectorAll('#view-tabs-nav .view-tab-btn')).map(b => b.getAttribute('data-view'))
    })""")

    print(f"  • Active View after loading: {hiding_status['activeView']} (Expected: explorer fallback)")
    print(f"  • Capabilities: changes={hiding_status['changesCap']}, definitions={hiding_status['defsCap']}")
    print(f"  • Available tabs: {hiding_status['availableTabs']}")
    print(f"  • Matrix tab in DOM: {hiding_status['matrixTabExists']} (Expected: False)")
    print(f"  • Glossary tab in DOM: {hiding_status['glossaryTabExists']} (Expected: False)")
    print(f"  • Disabled placeholder tabs: {hiding_status['disabledTabsCount']} (Expected: 0)")

    assert hiding_status['activeView'] == "explorer", f"Active view should fall back to explorer, got {hiding_status['activeView']}"
    assert hiding_status['changesCap'] is False, "capabilities.changes should be False"
    assert hiding_status['defsCap'] == 0, "capabilities.definitions should be 0"
    assert not hiding_status['matrixTabExists'], "Matrix tab should not be rendered when changes=False"
    assert not hiding_status['glossaryTabExists'], "Glossary tab should not be rendered when definitions=0"
    assert hiding_status['disabledTabsCount'] == 0, "Unavailable tabs should not be rendered as disabled placeholders"
    assert set(hiding_status['availableTabs']) == {"explorer", "analytics"}, "Only Explorer and Analytics tabs should be rendered"

    # 7. Deep Link Fallback Verification
    print("\n--- 7. Deep Link Fallback to Explorer ---")
    # Hash requesting unavailable view 'glossary'
    client.evaluate("location.hash = '#view=glossary'")
    time.sleep(0.2)
    client.evaluate("applyHash()")
    active_after_hash = client.evaluate("state.activeView")
    print(f"  • Active view after hash '#view=glossary': {active_after_hash} (Expected: explorer)")
    assert active_after_hash == "explorer", "Unavailable view hash should fall back to explorer"

    # Hash requesting non-existent element
    client.evaluate("location.hash = '#ref=99999999'")
    time.sleep(0.2)
    client.evaluate("applyHash()")
    active_after_ref = client.evaluate("state.activeView")
    print(f"  • Active view after invalid '#ref=99999999': {active_after_ref} (Expected: explorer)")
    assert active_after_ref == "explorer", "Unknown element ref should fall back safely to explorer without error"

    # Restore Healthcare Provider for remaining test suites
    print("\n--- 8. Reloading Healthcare Provider for Subsequent Suites ---")
    load_workbook(client, healthcare_path)
    final_count = client.evaluate("state.data ? state.data.elements.length : 0")
    print(f"  • Restored Healthcare elements: {final_count} (Expected: 1985)")
    assert final_count == 1985

    print("\n" + "="*70)
    print("ALL EMPTY START, STATE MACHINE & CAPABILITIES TESTS PASSED 100%!")
    print("="*70)
