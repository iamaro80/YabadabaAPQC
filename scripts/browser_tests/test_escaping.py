"""
HTML escaping, template sanitization, and XSS prevention browser tests.
"""

import os
import json
import time
import base64
import openpyxl
import shutil
from browser_tests.cdp_client import load_workbook

def run_escaping_tests(client, project_dir):
    print("\n" + "="*70)
    print("ITEM 3: WORKBOOK ESCAPING & INLINE HANDLER REMOVAL TEST")
    print("="*70)

    # 1. Existing Item 3: Malicious element ID and Name escaping
    malicious_xlsx = "/tmp/malicious_pcf.xlsx"
    src_xlsx = os.path.join(project_dir, "K07128_Healthcare Provider_v721_vs_Healthcare Provider_v611_April 2023.xlsx")
    shutil.copyfile(src_xlsx, malicious_xlsx)
    wb = openpyxl.load_workbook(malicious_xlsx)
    ws = wb['Combined']
    mal_id = '99999\'"><script>window.__xss_executed=true;</script>\\\\test'
    mal_name = 'XSS \'single\' "double" <script>window.__xss_executed=true;</script> \\\\backslash'
    ws.cell(row=2, column=1, value=mal_id)
    ws.cell(row=2, column=3, value=mal_name)
    wb.save(malicious_xlsx)

    with open(malicious_xlsx, "rb") as mf:
        m_b64 = base64.b64encode(mf.read()).decode("ascii")

    client.evaluate(f"""(() => {{
        const b64 = "{m_b64}";
        const bin = atob(b64);
        const bytes = new Uint8Array(bin.length);
        for (let i = 0; i < bin.length; i++) bytes[i] = bin.charCodeAt(i);
        pendingWorkbookBuffer = bytes;
        pendingFileName = "malicious_pcf.xlsx";
        openUploadModal(bytes, "malicious_pcf.xlsx");
        document.getElementById("confirm-upload-btn").click();
    }})()""")
    time.sleep(0.5)

    xss_executed = client.evaluate("Boolean(window.__xss_executed)")
    print(f"XSS executed after uploading workbook with quotes and script tags: {xss_executed}")
    assert not xss_executed, "XSS vulnerability detected! Script executed from workbook data."

    # Verify element links and text rendering
    raw_id_json = json.dumps(mal_id)
    jump_test = client.evaluate(f"""(() => {{
        const rawId = {raw_id_json};
        jumpToElement(rawId);
        const titleEl = document.querySelector('.ref-title');
        const breadcrumbEl = document.querySelector('.ref-breadcrumb');
        return {{
            titleText: titleEl ? titleEl.textContent : '',
            breadcrumbText: breadcrumbEl ? breadcrumbEl.textContent : '',
            xssFired: Boolean(window.__xss_executed)
        }};
    }})()""")
    print(f"Escaped title rendered safely: '{jump_test['titleText'][:50]}...'")
    print(f"XSS fired after jumping: {jump_test['xssFired']}")
    assert not jump_test['xssFired'], "XSS script executed during jumpToElement!"
    assert "XSS" in jump_test['titleText'], "Malicious element title failed to render"

    # Restore Healthcare PCF
    load_workbook(client, src_xlsx)
    assert client.evaluate("state.data.elements.length") == 1985

    # 2. Item 1 Security Test: Category ID Validation & Attribute Injection Prevention
    print("\n" + "="*70)
    print("ITEM 1 SECURITY TEST: MALICIOUS CATEGORY SEGMENT IN HIERARCHY ID & CHECKBOX ESCAPING")
    print("="*70)

    # Instrument window.alert and tracking flags
    client.evaluate("""(() => {
        window.__alert_called = false;
        window.__orig_alert = window.alert;
        window.alert = function(msg) {
            window.__alert_called = true;
            console.warn("ALERT INTERCEPTED:", msg);
        };
        window.__xss_category = false;
    })()""")

    malicious_cat_xlsx = "/tmp/malicious_cat_pcf.xlsx"
    shutil.copyfile(src_xlsx, malicious_cat_xlsx)
    wb_cat = openpyxl.load_workbook(malicious_cat_xlsx)
    ws_cat = wb_cat['Combined']

    # Inject malicious category segment in Hierarchy ID at row 2
    # e.g. 1" onfocus="alert(1)'<script>window.__xss_category=true;</script>\\.0
    mal_hid = '1" onfocus="alert(1)\'<script>window.__xss_category=true;</script>\\\\.1.0'
    ws_cat.cell(row=2, column=2, value=mal_hid)
    wb_cat.save(malicious_cat_xlsx)

    with open(malicious_cat_xlsx, "rb") as mcf:
        mc_b64 = base64.b64encode(mcf.read()).decode("ascii")

    # Upload workbook with malicious category segment
    upload_res = client.evaluate(f"""(() => {{
        const b64 = "{mc_b64}";
        const bin = atob(b64);
        const bytes = new Uint8Array(bin.length);
        for (let i = 0; i < bin.length; i++) bytes[i] = bin.charCodeAt(i);
        pendingWorkbookBuffer = bytes;
        pendingFileName = "malicious_cat_pcf.xlsx";
        openUploadModal(bytes, "malicious_cat_pcf.xlsx");
        document.getElementById("confirm-upload-btn").click();
        
        // Check if malicious hierarchy id was rejected
        const foundMalElem = state.data.elements.some(e => e.hierarchy_id === {json.dumps(mal_hid)});
        return {{
            rejected: !foundMalElem,
            totalElements: state.data.elements.length
        }};
    }})()""")
    time.sleep(0.5)

    print(f"Malicious category segment row rejected: {upload_res['rejected']} (Total elements: {upload_res['totalElements']})")
    assert upload_res['rejected'], "Parser accepted invalid category ID segment that does not match /^\\d+$/"

    # Also test upload dialog checkbox escaping directly with malicious category object
    checkbox_test = client.evaluate("""(() => {
        const dummyCategories = [{
            category_id: '1" onfocus="alert(1)\\'<script>window.__xss_category=true;</script>\\\\',
            category_name: 'Injected Category',
            hierarchy_id: '1.0'
        }];
        const catsGrid = document.getElementById('modal-op-cats-container');
        catsGrid.innerHTML = "";
        dummyCategories.forEach(c => {
            const label = document.createElement("label");
            label.className = "op-cat-checkbox-item";
            label.innerHTML = `
                <input type="checkbox" value="${escapeHtml(c.category_id)}" checked>
                <span style="font-weight:600; font-family:var(--font-mono);">${escapeHtml(c.hierarchy_id)}</span>
                <span style="color:var(--text-muted);">${escapeHtml(c.category_name)}</span>
            `;
            catsGrid.appendChild(label);
        });

        // Trigger focus on all inputs to test if any inline onfocus attribute would fire
        const inputs = Array.from(catsGrid.querySelectorAll('input'));
        inputs.forEach(input => input.focus());

        const hasOnfocusAttr = catsGrid.querySelector('[onfocus]') !== null;
        const injectedVal = inputs[0] ? inputs[0].value : '';

        return {
            hasOnfocusAttr,
            injectedVal,
            alertCalled: Boolean(window.__alert_called),
            xssCategory: Boolean(window.__xss_category)
        };
    })()""")

    print(f"Checkbox injection test: hasOnfocusAttr={checkbox_test['hasOnfocusAttr']}, alertCalled={checkbox_test['alertCalled']}, xssCategory={checkbox_test['xssCategory']}")
    print(f"Safely preserved attribute value: {checkbox_test['injectedVal']}")
    assert not checkbox_test['hasOnfocusAttr'], "HTML attribute injection detected in upload categories grid!"
    assert not checkbox_test['alertCalled'], "window.alert was called via attribute injection!"
    assert not checkbox_test['xssCategory'], "Script tag in category ID was executed!"

    # Restore default PCF and clean up alert mock
    client.evaluate("""(() => {
        if (window.__orig_alert) window.alert = window.__orig_alert;
        delete window.__alert_called;
        delete window.__xss_category;
    })()""")
    load_workbook(client, src_xlsx)
    assert client.evaluate("state.data.elements.length") == 1985
    print("Category ID security verification passed completely.")
